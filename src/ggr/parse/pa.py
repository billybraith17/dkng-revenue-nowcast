"""Parser for PGCB monthly sports wagering report summaries.

Layout: licensee blocks down column A, months across columns. Each block has
sections (Total / Retail / Online Sports Wagering) whose metric rows follow.
Only the Online section is emitted, and PA's metric set maps one-to-one onto
the schema's promo columns — the reason PA is valuable:

    Handle*                  → handle
    Revenue                  → gross_revenue   (pre-promo GGR)
    Promotional Credits      → promo_credits
    Gross Revenue (Taxable)  → adjusted_revenue

Trust nothing but the month header cells ("July 2022"): sheet names lie (the
FY2022-2023 file's data sheet is titled "FY 2021-22"), and fiscal-year files
span two calendar years.
"""

from __future__ import annotations

import datetime as dt
import re
from calendar import monthrange
from pathlib import Path

import openpyxl
import pandas as pd

from ggr.operators import canonicalize
from ggr.schema import tidy_frame

MONTH_HEADER = re.compile(
    r"^(january|february|march|april|may|june|july|august|september|october|november|december)"
    r"\s+(\d{4})$"
)
MONTH_NUM = {m: i + 1 for i, m in enumerate(
    "january february march april may june july august september october november december".split()
)}

SECTIONS = {"total sports wagering", "retail sports wagering", "online sports wagering"}
METRICS = {
    "handle*": "handle",
    "revenue": "gross_revenue",
    "promotional credits": "promo_credits",
    "gross revenue (taxable)": "adjusted_revenue",
}
IGNORED_PREFIXES = ("state tax", "local share", "monthly sports wagering", "footnotes")
NOT_A_LICENSEE = {"grand total"}


def _norm(cell) -> str:
    return re.sub(r"\s+", " ", str(cell)).strip().lower() if cell is not None else ""


def _month_columns(rows: list[tuple]) -> dict[int, tuple[int, int]]:
    """Map column index → (year, month) from the header row."""
    for row in rows:
        cols = {}
        for j, cell in enumerate(row):
            m = MONTH_HEADER.match(_norm(cell))
            if m:
                cols[j] = (int(m.group(2)), MONTH_NUM[m.group(1)])
        if len(cols) >= 2:
            return cols
    raise ValueError("no month header row found")


def parse_workbook(path: Path, source_url: str) -> pd.DataFrame:
    fetched_at = dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records: dict[tuple[str, tuple[int, int]], dict] = {}
    for sheet in wb.sheetnames:
        if _norm(sheet) == "footnotes":
            continue
        rows = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
        month_cols = _month_columns(rows)
        licensee, section = None, None
        for row in rows:
            label = _norm(row[0])
            if not label or label.startswith(IGNORED_PREFIXES):
                continue
            if label in SECTIONS:
                section = label
            elif label in METRICS:
                if licensee is None or section != "online sports wagering":
                    continue
                field = METRICS[label]
                for j, (year, month) in month_cols.items():
                    value = row[j] if j < len(row) else None
                    if value is None or (isinstance(value, str) and not value.strip()):
                        continue
                    key = (licensee, (year, month))
                    existing = records.setdefault(key, {})
                    if field in existing and existing[field] != float(value):
                        raise ValueError(
                            f"{path.name}: conflicting {field} for {licensee} "
                            f"{year}-{month:02d} — duplicated block with different values"
                        )
                    existing[field] = float(value)
            elif label in NOT_A_LICENSEE:
                licensee, section = None, None
            else:
                licensee, section = str(row[0]).strip(), None
    wb.close()

    tidy_rows = []
    for (licensee, (year, month)), fields in records.items():
        tidy_rows.append(
            {
                "state": "PA",
                "period_start": dt.date(year, month, 1),
                "period_end": dt.date(year, month, monthrange(year, month)[1]),
                "operator_raw": licensee,
                "operator_canonical": canonicalize(licensee),
                "product": "osb",
                "source_url": source_url,
                "fetched_at": fetched_at,
                **fields,
            }
        )
    df = tidy_frame(tidy_rows)
    unknown = df.loc[df["operator_canonical"].isna(), "operator_raw"].unique()
    if len(unknown):
        raise ValueError(f"{path.name}: unknown operators {sorted(unknown)} — add to ggr.operators")
    return df.sort_values(["period_end", "operator_canonical"]).reset_index(drop=True)
