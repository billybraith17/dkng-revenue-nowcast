"""Parser for Michigan MGCB internet gaming / internet sports betting workbooks.

Layout (one sheet per calendar year, current file also carries the prior year):
operators run across columns in fixed-width groups, months down rows.

    row "Platform Providers"  →  brand name at each group's first column
    row "Month"               →  metric labels; a group starts at "Total Handle"
                                 (sports) or "Gross Receipts" (iGaming)
    rows January..December    →  values; unpublished months are blank

Michigan is one of the few states reporting adjusted gross receipts (post-promo)
next to gross — both are kept. State tax columns are ignored.
"""

from __future__ import annotations

import datetime as dt
import re
from calendar import monthrange
from pathlib import Path

import openpyxl
import pandas as pd

from ggr.operators import canonicalize
from ggr.schema import tidy_frame

MONTHS = {
    m: i + 1
    for i, m in enumerate(
        "january february march april may june july august september october november december".split()
    )
}

GROUP_SPECS = {
    "osb": {
        "start": "total handle",
        "fields": {0: "handle", 1: "gross_revenue", 2: "adjusted_revenue"},
    },
    "igaming": {
        "start": "gross receipts",
        "fields": {0: "gross_revenue", 1: "adjusted_revenue"},
    },
}


def _norm(cell) -> str:
    return re.sub(r"\s+", " ", str(cell)).strip().lower() if cell is not None else ""


def _number(cell) -> float | None:
    if cell is None or isinstance(cell, (int, float)):
        return cell
    text = str(cell).replace("$", "").replace(",", "").strip().lower()
    if text in ("", "-", "*", "n/a") or text.startswith("#"):  # '#REF!' etc. occur in the wild
        return None
    return float(text)


def _find_row(rows: list[tuple], prefix: str) -> tuple | None:
    for row in rows:
        if _norm(row[0]).startswith(prefix):
            return row
    return None


def parse_workbook(path: Path, product: str, source_url: str) -> pd.DataFrame:
    spec = GROUP_SPECS[product]
    fetched_at = dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet in wb.sheetnames:
        year_text = sheet.strip()[-4:]
        if not year_text.isdigit():
            continue
        year = int(year_text)
        rows = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
        providers = _find_row(rows, "platform provider")
        labels = _find_row(rows, "month")
        if providers is None or labels is None:
            raise ValueError(f"{path.name}[{sheet}]: header rows not found")
        starts = [j for j, cell in enumerate(labels) if _norm(cell) == spec["start"]]
        if not starts:
            raise ValueError(f"{path.name}[{sheet}]: no {spec['start']!r} columns found")

        for row in rows:
            month = MONTHS.get(_norm(row[0]))
            if month is None:
                continue
            for j in starts:
                operator_raw = str(providers[j]).strip() if j < len(providers) and providers[j] else ""
                if not operator_raw or "total" in operator_raw.lower():
                    continue
                values = {
                    field: _number(row[j + off]) if j + off < len(row) else None
                    for off, field in spec["fields"].items()
                }
                if all(v is None for v in values.values()):
                    continue
                records.append(
                    {
                        "state": "MI",
                        "period_start": dt.date(year, month, 1),
                        "period_end": dt.date(year, month, monthrange(year, month)[1]),
                        "operator_raw": operator_raw,
                        "operator_canonical": canonicalize(operator_raw),
                        "product": product,
                        "source_url": source_url,
                        "fetched_at": fetched_at,
                        **values,
                    }
                )
    wb.close()
    df = tidy_frame(records)
    unknown = df.loc[df["operator_canonical"].isna(), "operator_raw"].unique()
    if len(unknown):
        raise ValueError(f"{path.name}: unknown operators {sorted(unknown)} — add to ggr.operators")
    return df.sort_values(["period_end", "operator_canonical"]).reset_index(drop=True)
