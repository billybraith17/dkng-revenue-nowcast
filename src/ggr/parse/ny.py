"""Parser for NY weekly mobile sports wagering workbooks.

Layout quirks the parser must survive (seen in real files, frozen in tests):
- one sheet per NY fiscal year (April–March), named like "FY 23-24";
- header row ("Week-Ending" / "Handle" / "GGR") sits at a different row per sheet,
  and the Handle/GGR *column positions differ between sheets*;
- date cells below the first are formulas (=A13+7), so cached values are required
  (``data_only=True``);
- future weeks are pre-filled with dates but empty values — skipped;
- GGR can be negative in a bad-hold week; negatives are kept.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pandas as pd

from ggr.operators import canonicalize
from ggr.schema import tidy_frame


def _header_columns(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """Find the header row index and the column index of Handle and GGR."""
    for i, row in enumerate(rows):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if cells and cells[0] == "week-ending":
            cols = {}
            for j, cell in enumerate(cells):
                if cell in ("handle", "ggr"):
                    cols[cell] = j
            if {"handle", "ggr"} <= cols.keys():
                return i, cols
    return None


def parse_workbook(path: Path, operator_raw: str, source_url: str) -> pd.DataFrame:
    fetched_at = dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet in wb.sheetnames:
        rows = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
        header = _header_columns(rows)
        if header is None:
            raise ValueError(f"{path.name}[{sheet}]: no Week-Ending/Handle/GGR header found")
        start, cols = header
        for row in rows[start + 1 :]:
            week_ending = row[0]
            if not isinstance(week_ending, dt.datetime):
                continue
            handle = row[cols["handle"]] if cols["handle"] < len(row) else None
            ggr = row[cols["ggr"]] if cols["ggr"] < len(row) else None
            if handle is None and ggr is None:
                continue
            records.append(
                {
                    "state": "NY",
                    "period_start": week_ending - dt.timedelta(days=6),
                    "period_end": week_ending,
                    "operator_raw": operator_raw,
                    "operator_canonical": canonicalize(operator_raw),
                    "product": "osb",
                    "handle": handle,
                    "gross_revenue": ggr,
                    "source_url": source_url,
                    "fetched_at": fetched_at,
                }
            )
    wb.close()
    df = tidy_frame(records)
    if df["operator_canonical"].isna().any():
        raise ValueError(f"{path.name}: unknown operator {operator_raw!r} — add to ggr.operators")
    df = _drop_zero_duplicates(df, path.name)
    return df.sort_values("period_end").reset_index(drop=True)


def _drop_zero_duplicates(df: pd.DataFrame, name: str) -> pd.DataFrame:
    """Some workbooks (Wynn) repeat a fiscal-year block with all-zero values.

    For a duplicated week, keep the row with the largest activity; refuse to
    guess if two conflicting non-zero rows exist.
    """
    activity = df["handle"].fillna(0).abs() + df["gross_revenue"].fillna(0).abs()
    df = df.assign(_activity=activity).sort_values("_activity", ascending=False)
    kept = df.drop_duplicates("period_end", keep="first")
    dropped = df.loc[df.index.difference(kept.index)]
    if (dropped["_activity"] != 0).any():
        raise ValueError(f"{name}: conflicting non-zero rows for the same week")
    return kept.drop(columns="_activity")
