"""Parser for Indiana Gaming Commission monthly revenue workbooks (sports).

The "Detail of Sports Wagering Tax" sheet is the prize: unlike NJ/PA, Indiana
prints **per-skin online detail** — rows like "AS - Sportsbook.DraftKings.com"
with their own handle and gross receipts, under each casino licensee.

Layout: two side-by-side column groups (NORTHERN / SOUTHERN licensees). Each
casino block starts with a header row (casino name, "Handle", …, "Gross
Receipts") followed by skin rows ("XX - domain"), a Retail row (skipped —
online-only scope), Adjustments and Taxable AGR rows (both skipped). Column
offsets differ between groups and eras, so handle/gross positions are read
from each block's own header row.
"""

from __future__ import annotations

import datetime as dt
import re
from calendar import monthrange
from pathlib import Path

import openpyxl
import pandas as pd

from ggr.operators import canonicalize
from ggr.schema import tidy_frame

DETAIL_SHEET_MARKER = "detail of sports wagering"
SKIN_ROW = re.compile(r"^[A-Z]{2,4}\s*-\s*\S")
NOT_SKIN_LABELS = {"retail", "adjustments", "taxable agr", "handle"}


def _norm(cell) -> str:
    return re.sub(r"\s+", " ", str(cell)).strip().lower() if cell is not None else ""


def _number(cell) -> float | None:
    if cell is None or isinstance(cell, (int, float)):
        return float(cell) if cell is not None else None
    text = str(cell).replace("$", "").replace(",", "").strip()
    if text in ("", "-"):
        return None
    if text.startswith("(") and text.endswith(")"):
        return -float(text[1:-1])
    return float(text)


def _find_detail_sheet(wb) -> list[tuple]:
    for name in wb.sheetnames:
        rows = [tuple(r) for r in wb[name].iter_rows(values_only=True)]
        head = " ".join(_norm(c) for row in rows[:4] for c in row)
        if DETAIL_SHEET_MARKER in head:
            return rows
    raise ValueError("no 'Detail of Sports Wagering' sheet found")


def parse_workbook(path: Path, year: int, month: int, source_url: str) -> pd.DataFrame:
    fetched_at = dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = _find_detail_sheet(wb)
    wb.close()

    month_name = dt.date(year, month, 1).strftime("%B").lower()
    sheet_text = " ".join(_norm(c) for row in rows[:4] for c in row)
    if month_name not in sheet_text or str(year) not in sheet_text:
        raise ValueError(
            f"{path.name}: detail sheet header {sheet_text[:100]!r} doesn't state "
            f"{month_name} {year}"
        )

    # header rows: casino name in some column j, "Handle" at j+1
    blocks = {}  # base col → (handle_offset, gross_offset), refreshed per header row
    records = []
    for row in rows:
        for j, cell in enumerate(row[:-1]):
            if _norm(row[j + 1]) == "handle" and _norm(cell) and _norm(cell) != "handle":
                gross = next(
                    (k for k in range(j + 2, min(j + 6, len(row))) if "gross" in _norm(row[k])),
                    None,
                )
                if gross is None:
                    raise ValueError(f"{path.name}: block header without Gross Receipts near {cell!r}")
                blocks[j] = (1, gross - j)
        for j, (h_off, g_off) in blocks.items():
            label = str(row[j]).strip() if j < len(row) and row[j] is not None else ""
            if not SKIN_ROW.match(label) or _norm(label) in NOT_SKIN_LABELS:
                continue
            handle = _number(row[j + h_off]) if j + h_off < len(row) else None
            gross_rev = _number(row[j + g_off]) if j + g_off < len(row) else None
            if handle is None and gross_rev is None:
                continue
            records.append(
                {
                    "state": "IN",
                    "period_start": dt.date(year, month, 1),
                    "period_end": dt.date(year, month, monthrange(year, month)[1]),
                    "operator_raw": label,
                    "operator_canonical": canonicalize(label),
                    "product": "osb",
                    "handle": handle,
                    "gross_revenue": gross_rev,
                    "source_url": source_url,
                    "fetched_at": fetched_at,
                }
            )

    if not records:
        raise ValueError(f"{path.name}: no online skin rows found in detail sheet")
    df = tidy_frame(records)
    unknown = df.loc[df["operator_canonical"].isna(), "operator_raw"].unique()
    if len(unknown):
        raise ValueError(f"{path.name}: unknown operators {sorted(unknown)} — add to ggr.operators")
    return df.sort_values("operator_raw").reset_index(drop=True)
