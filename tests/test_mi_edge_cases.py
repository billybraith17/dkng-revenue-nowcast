"""MI parser edge cases: number cleaning, group walking, provider skipping,
unknown operators, and the igaming vs osb column-group specs.
"""

import openpyxl
import pytest

from ggr.parse.mi import _number, parse_workbook


def _make_osb_workbook(tmp_path, provider="BetMGM", month_rows=None, title="Internet Sports Betting 2024"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["", "2024 Internet Sports Betting Revenue"])
    ws.append(["Operators", "MGM Grand Detroit", "", "", "", "Statewide"])
    ws.append(["Casino Name", "MGM Grand Detroit", "", "", "", ""])
    ws.append(["Platform Providers", provider, "", "", "", "Total"])
    ws.append(["Initial date of Operation launch", "2021-01-22", "", "", "", ""])
    ws.append(
        [
            "Month",
            "Total Handle",
            "Gross Sports Betting Receipts",
            "Adjusted Gross Sports Betting Receipts",
            "Internet Sports Betting State Tax",
            "Total Handle",  # the "Total" column group — must be skipped
        ]
    )
    for row in month_rows or []:
        ws.append(row)
    path = tmp_path / "mi_synthetic.xlsx"
    wb.save(path)
    return path


class TestNumberCleaning:
    def test_plain_and_negative(self):
        assert _number(123.45) == 123.45
        assert _number(-50) == -50

    def test_currency_strings(self):
        assert _number("$1,234.56") == pytest.approx(1234.56)

    def test_null_markers(self):
        for marker in (None, "", "-", "*", "n/a", "N/A"):
            assert _number(marker) is None

    def test_excel_errors_are_null(self):
        for err in ("#REF!", "#VALUE!", "#DIV/0!", "#N/A"):
            assert _number(err) is None

    def test_garbage_still_raises(self):
        with pytest.raises(ValueError):
            _number("twelve")


def test_month_rows_parsed_and_total_group_skipped(tmp_path):
    path = _make_osb_workbook(
        tmp_path,
        month_rows=[
            ["January", 1000.0, 100.0, 80.0, 5.0, 99999.0],
            ["February", "", "", "", "", ""],  # unpublished month
            ["Total", 1000.0, 100.0, 80.0, 5.0, 99999.0],  # summary row
        ],
    )
    df = parse_workbook(path, "osb", "https://example.test")
    assert len(df) == 1  # one operator-month; Total column and Total row both skipped
    row = df.iloc[0]
    assert row["operator_canonical"] == "betmgm"
    assert row["handle"] == 1000.0
    assert row["gross_revenue"] == 100.0
    assert row["adjusted_revenue"] == 80.0
    assert str(row["period_start"].date()) == "2024-01-01"
    assert str(row["period_end"].date()) == "2024-01-31"


def test_ref_error_cell_becomes_null_not_crash(tmp_path):
    path = _make_osb_workbook(tmp_path, month_rows=[["March", "#REF!", 100.0, 80.0, 5.0, ""]])
    df = parse_workbook(path, "osb", "https://example.test")
    assert df["handle"].isna().all()
    assert df["gross_revenue"].iloc[0] == 100.0


def test_unknown_operator_raises(tmp_path):
    path = _make_osb_workbook(
        tmp_path, provider="Mystery Books", month_rows=[["January", 1.0, 1.0, 1.0, 0.0, ""]]
    )
    with pytest.raises(ValueError, match="unknown operators.*Mystery Books"):
        parse_workbook(path, "osb", "https://example.test")


def test_non_year_sheet_ignored(tmp_path):
    path = _make_osb_workbook(
        tmp_path, month_rows=[["January", 1.0, 1.0, 1.0, 0.0, ""]], title="Notes"
    )
    # Only sheet has no trailing year → nothing parsed → empty frame is an error
    # in build(), but parse itself returns an empty schema frame.
    df = parse_workbook(path, "osb", "https://example.test")
    assert df.empty


def test_igaming_three_column_groups(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Internet Gaming 2024"
    ws.append(["", "2024 Internet Gaming Revenue"])
    ws.append(["Operator", "MGM Grand Detroit", "", "", "MotorCity Casino"])
    ws.append(["Casino Name", "MGM Grand Detroit", "", "", "MotorCity Casino"])
    ws.append(["Platform Provider", " BetMGM", "", "", "FanDuel"])
    ws.append(["Initial date of Operation launch", "2021-01-22", "", "", "2021-01-22"])
    # note the double spaces, as printed in the real files
    ws.append(["Month", "Gross  Receipts", "Adjusted Gross  Receipts", "Internet Gaming State Tax", "Gross  Receipts", "Adjusted Gross  Receipts", "Internet Gaming State Tax"])
    ws.append(["January", 500.0, 450.0, 90.0, 700.0, 650.0, 130.0])
    path = tmp_path / "mi_ig_synthetic.xlsx"
    wb.save(path)

    df = parse_workbook(path, "igaming", "https://example.test")
    assert len(df) == 2
    by_op = df.set_index("operator_canonical")
    assert by_op.loc["betmgm", "gross_revenue"] == 500.0
    assert by_op.loc["betmgm", "adjusted_revenue"] == 450.0
    assert by_op.loc["fanduel", "gross_revenue"] == 700.0
    assert df["handle"].isna().all()  # igaming has no handle
    assert (df["product"] == "igaming").all()


def test_missing_header_rows_raise(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Internet Sports Betting 2024"
    ws.append(["nothing useful here"])
    path = tmp_path / "mi_broken.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="header rows not found"):
        parse_workbook(path, "osb", "https://example.test")
